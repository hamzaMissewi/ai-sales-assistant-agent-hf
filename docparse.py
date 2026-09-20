"""Extract searchable text from common file types for the RAG knowledge base.

Supported: markdown/txt, csv/tsv, xlsx, pdf, docx. Images (png/jpg/etc.) are
OCR'd locally with tesseract when available; otherwise they are stored but
reported as not indexed.
"""

import csv
import io
import shutil
from pathlib import Path

TEXT_EXT = {".md", ".markdown", ".txt", ".text"}
CSV_EXT = {".csv", ".tsv"}
SHEET_EXT = {".xlsx"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}


def extract_text(path: Path) -> str:
    """Return the plain-text content of a file, or '' when not parseable."""
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        return _read_text(path)
    if ext in CSV_EXT:
        return _read_csv(path, "\t" if ext == ".tsv" else ",")
    if ext in SHEET_EXT:
        return _read_xlsx(path)
    if ext in PDF_EXT:
        return _read_pdf(path)
    if ext in DOCX_EXT:
        return _read_docx(path)
    if ext in IMG_EXT:
        return _read_ocr(path)
    return ""


def _read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")
    except OSError:
        return ""


def _read_csv(path: Path, delim: str) -> str:
    try:
        rows = list(
            csv.reader(io.StringIO(path.read_text(encoding="utf-8")), delimiter=delim)
        )
    except Exception:  # noqa: BLE001 - unreadable file -> not indexed
        return ""
    cells = []
    for row in rows:
        vals = [c.strip() for c in row if c and c.strip()]
        if vals:
            cells.append(" | ".join(vals))
    return "\n".join(cells)


def _read_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - corrupt file -> not indexed
        return ""
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(str(path))
        return "\n\n".join(p.extract_text() or "" for p in reader.pages)
    except Exception:  # noqa: BLE001 - scanned/corrupt pdf -> not indexed
        return ""


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    try:
        doc = Document(str(path))
    except Exception:  # noqa: BLE001 - corrupt docx -> not indexed
        return ""
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _read_ocr(path: Path) -> str:
    if shutil.which("tesseract") is None:
        return ""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        img = Image.open(path)
        return pytesseract.image_to_string(img)
    except Exception:  # noqa: BLE001 - OCR failure -> not indexed
        return ""
